import re
import unicodedata
from collections import Counter
from io import BytesIO


EXCEL_ERROR_VALUES = {
    "#REF!",
    "#N/A",
    "#VALUE!",
    "#NAME?",
    "#DIV/0!",
    "#NUM!",
    "#NULL!",
}


def normalize_header(value):
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]", "", text.upper())


def clean_value(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() in EXCEL_ERROR_VALUES:
        return None
    return text


def valid_isin(value):
    text = clean_value(value)
    if text is None:
        return None
    text = text.upper().replace(" ", "")
    if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", text):
        return text
    return None


def get_cell(row, header_indexes, header_name):
    index = header_indexes.get(header_name)
    if index is None or index >= len(row):
        return None
    return clean_value(row[index])


def load_mgc_workbook(file_content):
    from openpyxl import load_workbook

    return load_workbook(
        filename=BytesIO(file_content),
        read_only=True,
        data_only=True,
    )


def extract_mgc_rows(file_content):
    workbook = load_mgc_workbook(file_content)

    selected_sheet = None
    selected_header_row = None
    selected_indexes = None
    required_headers = {"NEMO", "NOMBRE", "ISIN", "EMISOR"}

    for sheet in workbook.worksheets:
        max_header_row = min(sheet.max_row or 15, 15)
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=max_header_row, values_only=True),
            start=1,
        ):
            indexes = {}
            for index, value in enumerate(row):
                header = normalize_header(value)
                if header and header not in indexes:
                    indexes[header] = index

            if required_headers.issubset(indexes):
                selected_sheet = sheet
                selected_header_row = row_number
                selected_indexes = indexes
                break

        if selected_sheet is not None:
            break

    if selected_sheet is None:
        raise ValueError(
            "No se encontró una hoja con las columnas NEMO, NOMBRE, ISIN y EMISOR."
        )

    rows = []
    for row_number, row in enumerate(
        selected_sheet.iter_rows(
            min_row=selected_header_row + 1,
            values_only=True,
        ),
        start=selected_header_row + 1,
    ):
        nemo = get_cell(row, selected_indexes, "NEMO")
        nombre = get_cell(row, selected_indexes, "NOMBRE")
        emisor = get_cell(row, selected_indexes, "EMISOR")

        if nemo is None and nombre is None and emisor is None:
            continue

        raw_isin = get_cell(row, selected_indexes, "ISIN")
        description = (
            get_cell(row, selected_indexes, "NUEVADESCRIPCION")
            or get_cell(row, selected_indexes, "DESCRIPCION")
        )

        rows.append(
            {
                "row_number": row_number,
                "nemo": nemo.upper() if nemo else None,
                "name": nombre,
                "isin": valid_isin(raw_isin),
                "raw_isin": raw_isin,
                "issuer": emisor,
                "issuer_id": get_cell(row, selected_indexes, "RUTISINEIN"),
                "description": description,
                "investor_url": get_cell(
                    row, selected_indexes, "LINKRELACIONINVERSIONISTA"
                ),
                "regulator_url": get_cell(row, selected_indexes, "LINKREGULADOR"),
                "underlying_market_source": get_cell(
                    row, selected_indexes, "MERCADOSUBYACENTE"
                ),
                "security_type_source": get_cell(
                    row, selected_indexes, "TIPODEVALOR"
                ),
                "country": get_cell(row, selected_indexes, "PAIS"),
                "sponsor": get_cell(row, selected_indexes, "PATROCINADOR"),
                "sponsor_nit": get_cell(row, selected_indexes, "NITPATROCINADOR"),
            }
        )

    nemo_counts = Counter(row["nemo"] for row in rows if row["nemo"])
    for row in rows:
        errors = []
        warnings = []

        if not row["nemo"]:
            errors.append("NEMO vacío")
        if not row["name"]:
            errors.append("Nombre vacío")
        if not row["issuer"]:
            errors.append("Emisor vacío")
        if row["nemo"] and nemo_counts[row["nemo"]] > 1:
            errors.append("NEMO duplicado")
        if row["raw_isin"] and row["isin"] is None:
            warnings.append("ISIN descartado por no ser válido")

        row["is_valid"] = len(errors) == 0
        row["validation_message"] = "; ".join(errors + warnings) or None

    return rows, selected_sheet.title
