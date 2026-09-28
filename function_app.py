import hashlib
import json
import logging
import os
import re
import unicodedata
from collections import Counter
from io import BytesIO

import azure.functions as func
import mssql_python
import requests
from openpyxl import load_workbook


app = func.FunctionApp()

EXCEL_ERROR_VALUES = {
    "#REF!",
    "#N/A",
    "#VALUE!",
    "#NAME?",
    "#DIV/0!",
    "#NUM!",
    "#NULL!",
}


def normalize_header(value):
    if value is None:
        return ""
    text = unicodedata.normalize("NFKD", str(value))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^A-Z0-9]", "", text.upper())


def clean_value(value):
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() in EXCEL_ERROR_VALUES:
        return None
    return text


def valid_isin(value):
    text = clean_value(value)
    if text is None:
        return None
    text = text.upper().replace(" ", "")
    if re.fullmatch(r"[A-Z]{2}[A-Z0-9]{9}[0-9]", text):
        return text
    return None


def get_cell(row, header_indexes, header_name):
    index = header_indexes.get(header_name)
    if index is None or index >= len(row):
        return None
    return clean_value(row[index])


def extract_mgc_rows(file_content):
    workbook = load_workbook(
        filename=BytesIO(file_content),
        read_only=True,
        data_only=True,
    )

    selected_sheet = None
    selected_header_row = None
    selected_indexes = None
    required_headers = {"NEMO", "NOMBRE", "ISIN", "EMISOR"}

    for sheet in workbook.worksheets:
        max_header_row = min(sheet.max_row or 15, 15)
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=1, max_row=max_header_row, values_only=True),
            start=1,
        ):
            indexes = {}
            for index, value in enumerate(row):
                header = normalize_header(value)
                if header and header not in indexes:
                    indexes[header] = index

            if required_headers.issubset(indexes):
                selected_sheet = sheet
                selected_header_row = row_number
                selected_indexes = indexes
                break

        if selected_sheet is not None:
            break

    if selected_sheet is None:
        raise ValueError(
            "No se encontró una hoja con las columnas NEMO, NOMBRE, ISIN y EMISOR."
        )

    rows = []
    for row_number, row in enumerate(
        selected_sheet.iter_rows(
            min_row=selected_header_row + 1,
            values_only=True,
        ),
        start=selected_header_row + 1,
    ):
        nemo = get_cell(row, selected_indexes, "NEMO")
        nombre = get_cell(row, selected_indexes, "NOMBRE")
        emisor = get_cell(row, selected_indexes, "EMISOR")

        if nemo is None and nombre is None and emisor is None:
            continue

        raw_isin = get_cell(row, selected_indexes, "ISIN")
        description = (
            get_cell(row, selected_indexes, "NUEVADESCRIPCION")
            or get_cell(row, selected_indexes, "DESCRIPCION")
        )

        rows.append(
            {
                "row_number": row_number,
                "nemo": nemo.upper() if nemo else None,
                "name": nombre,
                "isin": valid_isin(raw_isin),
                "raw_isin": raw_isin,
                "issuer": emisor,
                "issuer_id": get_cell(row, selected_indexes, "RUTISINEIN"),
                "description": description,
                "investor_url": get_cell(
                    row, selected_indexes, "LINKRELACIONINVERSIONISTA"
                ),
                "regulator_url": get_cell(row, selected_indexes, "LINKREGULADOR"),
                "underlying_market_source": get_cell(
                    row, selected_indexes, "MERCADOSUBYACENTE"
                ),
                "security_type_source": get_cell(
                    row, selected_indexes, "TIPODEVALOR"
                ),
                "country": get_cell(row, selected_indexes, "PAIS"),
                "sponsor": get_cell(row, selected_indexes, "PATROCINADOR"),
                "sponsor_nit": get_cell(row, selected_indexes, "NITPATROCINADOR"),
            }
        )

    nemo_counts = Counter(row["nemo"] for row in rows if row["nemo"])
    for row in rows:
        errors = []
        warnings = []

        if not row["nemo"]:
            errors.append("NEMO vacío")
        if not row["name"]:
            errors.append("Nombre vacío")
        if not row["issuer"]:
            errors.append("Emisor vacío")
        if row["nemo"] and nemo_counts[row["nemo"]] > 1:
            errors.append("NEMO duplicado")
        if row["raw_isin"] and row["isin"] is None:
            warnings.append("ISIN descartado por no ser válido")

        row["is_valid"] = len(errors) == 0
        row["validation_message"] = "; ".join(errors + warnings) or None

    return rows, selected_sheet.title


def get_filename(response):
    disposition = response.headers.get("Content-Disposition", "")
    match = re.search(r'filename\*?=(?:UTF-8\'\')?"?([^";]+)', disposition)
    if match:
        return match.group(1).strip()
    return "Valores_listados_MGC.xlsx"


def insert_stage_rows(cursor, load_id, rows):
    statement = """
        INSERT INTO dbo.BvcMgcValorStage
        (
            CargaID, FilaOrigen, Nemo, Nombre, Isin, Emisor,
            IdentificadorEmisor, Descripcion, UrlRelacionInversionista,
            UrlRegulador, MercadoSubyacenteFuente, TipoValorFuente,
            Pais, Patrocinador, NitPatrocinador, EsValida, ErrorValidacion
        )
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?);
    """

    parameters = [
        (
            load_id,
            row["row_number"],
            row["nemo"],
            row["name"],
            row["isin"],
            row["issuer"],
            row["issuer_id"],
            row["description"],
            row["investor_url"],
            row["regulator_url"],
            row["underlying_market_source"],
            row["security_type_source"],
            row["country"],
            row["sponsor"],
            row["sponsor_nit"],
            1 if row["is_valid"] else 0,
            row["validation_message"],
        )
        for row in rows
    ]
    cursor.executemany(statement, parameters)


def upsert_catalog(cursor, load_id, rows):
    cursor.execute(
        """
        UPDATE catalog
           SET AusenciasConsecutivas = AusenciasConsecutivas + 1,
               Activo = CASE
                   WHEN AusenciasConsecutivas + 1 >= 2 THEN 0
                   ELSE Activo
               END,
               FechaActualizacionUTC = SYSUTCDATETIME()
        FROM dbo.BvcMgcValor AS catalog
        WHERE NOT EXISTS
        (
            SELECT 1
            FROM dbo.BvcMgcValorStage AS stage
            WHERE stage.CargaID = ?
              AND stage.EsValida = 1
              AND stage.Nemo = catalog.Nemo
        );
        """,
        load_id,
    )

    valid_rows = [row for row in rows if row["is_valid"]]
    for row in valid_rows:
        cursor.execute(
            "SELECT BvcMgcValorID FROM dbo.BvcMgcValor WHERE Nemo = ?;",
            row["nemo"],
        )
        existing = cursor.fetchone()

        common_parameters = (
            row["name"],
            row["isin"],
            row["issuer"],
            row["issuer_id"],
            row["description"],
            row["investor_url"],
            row["regulator_url"],
            row["underlying_market_source"],
            row["security_type_source"],
            row["country"],
            row["sponsor"],
            row["sponsor_nit"],
        )

        if existing:
            cursor.execute(
                """
                UPDATE dbo.BvcMgcValor
                   SET Nombre = ?,
                       Isin = ?,
                       Emisor = ?,
                       IdentificadorEmisor = ?,
                       Descripcion = ?,
                       UrlRelacionInversionista = ?,
                       UrlRegulador = ?,
                       MercadoSubyacenteFuente = ?,
                       TipoValorFuente = ?,
                       Pais = ?,
                       Patrocinador = ?,
                       NitPatrocinador = ?,
                       UltimaCargaID = ?,
                       AusenciasConsecutivas = 0,
                       Activo = 1,
                       FechaActualizacionUTC = SYSUTCDATETIME()
                 WHERE Nemo = ?;
                """,
                *common_parameters,
                load_id,
                row["nemo"],
            )
        else:
            cursor.execute(
                """
                INSERT INTO dbo.BvcMgcValor
                (
                    Nemo, Nombre, Isin, Emisor, IdentificadorEmisor,
                    Descripcion, UrlRelacionInversionista, UrlRegulador,
                    MercadoSubyacenteFuente, TipoValorFuente, Pais,
                    Patrocinador, NitPatrocinador, PrimeraCargaID,
                    UltimaCargaID, AusenciasConsecutivas, Activo
                )
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 0, 1);
                """,
                row["nemo"],
                *common_parameters,
                load_id,
                load_id,
            )


def run_bvc_load():
    page_url = os.environ["BVC_MGC_PAGE_URL"]
    file_url = os.environ["BVC_MGC_FILE_URL"]
    connection_string = os.environ["SQL_CONNECTION_STRING"]

    response = requests.get(
        file_url,
        headers={
            "User-Agent": "Mozilla/5.0 (compatible; BVC-MGC-Loader/1.0)",
            "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,*/*",
            "Referer": page_url,
        },
        timeout=(20, 120),
    )
    response.raise_for_status()

    file_content = response.content
    if len(file_content) < 1000 or not file_content.startswith(b"PK"):
        raise ValueError("La respuesta recibida no parece ser un archivo XLSX válido.")

    file_hash = hashlib.sha256(file_content).hexdigest()
    filename = get_filename(response)
    etag = response.headers.get("ETag")
    last_modified = response.headers.get("Last-Modified")
    rows, sheet_name = extract_mgc_rows(file_content)

    total_rows = len(rows)
    valid_count = sum(1 for row in rows if row["is_valid"])
    rejected_count = total_rows - valid_count

    connection = None
    load_id = None

    try:
        connection = mssql_python.connect(connection_string)
        cursor = connection.cursor()

        cursor.execute(
            """
            SELECT TOP (1) TotalFilas, FilasValidas, FilasRechazadas
            FROM dbo.BvcMgcCarga
            WHERE HashSha256 = ?
              AND Estado = 'OK'
            ORDER BY CargaID DESC;
            """,
            file_hash,
        )
        previous_same_file = cursor.fetchone()

        if previous_same_file:
            cursor.execute(
                """
                INSERT INTO dbo.BvcMgcCarga
                (
                    FuentePaginaUrl, FuenteArchivoUrl, NombreArchivo, ETag,
                    UltimaModificacionFuente, HashSha256, FechaFinUTC, Estado,
                    TotalFilas, FilasValidas, FilasRechazadas, Mensaje
                )
                VALUES (?, ?, ?, ?, ?, ?, SYSUTCDATETIME(), 'SIN_CAMBIOS', ?, ?, ?, ?);
                """,
                page_url,
                file_url,
                filename,
                etag,
                last_modified,
                file_hash,
                previous_same_file[0],
                previous_same_file[1],
                previous_same_file[2],
                "El hash SHA-256 ya había sido procesado correctamente.",
            )
            connection.commit()
            return {
                "status": "SIN_CAMBIOS",
                "hash": file_hash,
                "file": filename,
            }

        cursor.execute(
            """
            INSERT INTO dbo.BvcMgcCarga
            (
                FuentePaginaUrl, FuenteArchivoUrl, NombreArchivo, ETag,
                UltimaModificacionFuente, HashSha256, Estado,
                TotalFilas, FilasValidas, FilasRechazadas, Mensaje
            )
            OUTPUT INSERTED.CargaID
            VALUES (?, ?, ?, ?, ?, ?, 'PROCESANDO', ?, ?, ?, ?);
            """,
            page_url,
            file_url,
            filename,
            etag,
            last_modified,
            file_hash,
            total_rows,
            valid_count,
            rejected_count,
            f"Hoja detectada: {sheet_name}",
        )
        load_id = cursor.fetchone()[0]
        connection.commit()

        insert_stage_rows(cursor, load_id, rows)

        cursor.execute(
            """
            SELECT TOP (1) FilasValidas
            FROM dbo.BvcMgcCarga
            WHERE Estado = 'OK'
              AND CargaID <> ?
            ORDER BY CargaID DESC;
            """,
            load_id,
        )
        previous_load = cursor.fetchone()
        previous_valid_count = previous_load[0] if previous_load else None

        requires_review = valid_count == 0 or (
            previous_valid_count
            and valid_count < previous_valid_count * 0.80
        )

        if requires_review:
            cursor.execute(
                """
                UPDATE dbo.BvcMgcCarga
                   SET Estado = 'REQUIERE_REVISION',
                       FechaFinUTC = SYSUTCDATETIME(),
                       Mensaje = ?
                 WHERE CargaID = ?;
                """,
                "La cantidad de filas válidas es cero o cayó más del 20%.",
                load_id,
            )
            connection.commit()
            return {
                "status": "REQUIERE_REVISION",
                "load_id": load_id,
                "total": total_rows,
                "valid": valid_count,
            }

        upsert_catalog(cursor, load_id, rows)

        cursor.execute(
            """
            UPDATE dbo.BvcMgcCarga
               SET Estado = 'OK',
                   FechaFinUTC = SYSUTCDATETIME(),
                   Mensaje = ?
             WHERE CargaID = ?;
            """,
            f"Carga completada desde la hoja {sheet_name}.",
            load_id,
        )
        connection.commit()

        return {
            "status": "OK",
            "load_id": load_id,
            "file": filename,
            "sheet": sheet_name,
            "total": total_rows,
            "valid": valid_count,
            "rejected": rejected_count,
        }

    except Exception as error:
        if connection is not None:
            try:
                connection.rollback()
                if load_id is not None:
                    cursor = connection.cursor()
                    cursor.execute(
                        """
                        UPDATE dbo.BvcMgcCarga
                           SET Estado = 'ERROR',
                               FechaFinUTC = SYSUTCDATETIME(),
                               Mensaje = ?
                         WHERE CargaID = ?;
                        """,
                        str(error)[:2000],
                        load_id,
                    )
                    connection.commit()
            except Exception:
                logging.exception("No fue posible registrar el error en Azure SQL.")
        raise
    finally:
        if connection is not None:
            connection.close()


@app.function_name(name="CargarBvcMgcMensual")
@app.timer_trigger(
    schedule="%TIMER_SCHEDULE%",
    arg_name="my_timer",
    run_on_startup=False,
    use_monitor=True,
)
def cargar_bvc_mgc_mensual(my_timer: func.TimerRequest) -> None:
    if my_timer.past_due:
        logging.warning("La ejecución mensual de la BVC se inició con retraso.")

    result = run_bvc_load()
    logging.info("Resultado de carga BVC MGC: %s", json.dumps(result, ensure_ascii=False))


@app.function_name(name="CargarBvcMgcManual")
@app.route(
    route="bvc-mgc/cargar",
    methods=["POST"],
    auth_level=func.AuthLevel.FUNCTION,
)
def cargar_bvc_mgc_manual(req: func.HttpRequest) -> func.HttpResponse:
    try:
        result = run_bvc_load()
        return func.HttpResponse(
            json.dumps(result, ensure_ascii=False),
            status_code=200,
            mimetype="application/json",
        )
    except Exception as error:
        logging.exception("Falló la carga manual BVC MGC.")
        return func.HttpResponse(
            json.dumps({"status": "ERROR", "message": str(error)}, ensure_ascii=False),
            status_code=500,
            mimetype="application/json",
        )
